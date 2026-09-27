"""
Export service for Action Items to CSV and Excel format.
"""

from typing import List, Dict, Any
import io
import pandas as pd


class ExportService:
    """Handles exporting filtered action items to CSV and Excel."""

    EXPORT_COLUMNS = [
        ("action_title", "Action Item"),
        ("description", "Description"),
        ("assignee", "Assignee"),
        ("project_name", "Jira Project"),
        ("jira_project_key", "Project Key"),
        ("priority", "Priority"),
        ("due_date", "Due Date"),
        ("action_type", "Action Type"),
        ("status", "Status"),
        ("confidence_score", "Confidence"),
        ("source_excerpt", "Source Excerpt"),
        ("meeting_name", "Meeting Name"),
        ("created_at", "Date Extracted"),
    ]

    @classmethod
    def prepare_dataframe(cls, items: List[Dict[str, Any]]) -> pd.DataFrame:
        """Converts raw action items to a formatted presentation DataFrame."""
        if not items:
            df = pd.DataFrame(columns=[label for _, label in cls.EXPORT_COLUMNS])
            return df

        rows = []
        for item in items:
            row = {}
            for field_key, col_label in cls.EXPORT_COLUMNS:
                val = item.get(field_key)
                if val is None or val == "":
                    val = "—"
                elif field_key == "confidence_score" and isinstance(val, (int, float)):
                    val = f"{val * 100:.0f}%"
                row[col_label] = val
            rows.append(row)

        return pd.DataFrame(rows)

    @classmethod
    def to_csv(cls, items: List[Dict[str, Any]]) -> bytes:
        """Export action items to CSV bytes (with UTF-8 BOM for Microsoft Excel)."""
        df = cls.prepare_dataframe(items)
        csv_buffer = io.StringIO()
        df.to_csv(csv_buffer, index=False, encoding="utf-8")
        # Add UTF-8 BOM so Excel opens special characters correctly
        return ("\ufeff" + csv_buffer.getvalue()).encode("utf-8")

    @classmethod
    def to_excel(cls, items: List[Dict[str, Any]]) -> bytes:
        """Export action items to a styled Excel (.xlsx) workbook."""
        df = cls.prepare_dataframe(items)
        excel_buffer = io.BytesIO()

        with pd.ExcelWriter(excel_buffer, engine="openpyxl") as writer:
            df.to_excel(writer, sheet_name="Action Items", index=False)
            workbook = writer.book
            worksheet = writer.sheets["Action Items"]

            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

            # Header styling
            header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
            thin_border = Border(
                left=Side(style="thin", color="E2E8F0"),
                right=Side(style="thin", color="E2E8F0"),
                top=Side(style="thin", color="E2E8F0"),
                bottom=Side(style="thin", color="E2E8F0"),
            )

            for col_idx, col_name in enumerate(df.columns, 1):
                cell = worksheet.cell(row=1, column=col_idx)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

            # Auto-adjust column widths
            for col in worksheet.columns:
                max_len = max(len(str(cell.value or "")) for cell in col)
                col_letter = col[0].column_letter
                worksheet.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 50)

        excel_buffer.seek(0)
        return excel_buffer.getvalue()
